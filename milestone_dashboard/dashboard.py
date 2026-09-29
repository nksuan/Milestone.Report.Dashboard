from html import escape
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .report import MilestoneReport

_NAVY = RGBColor(20, 35, 61)
_BLUE = RGBColor(37, 99, 235)
_STATUS_COLORS = {
    "on track": "#16a34a",
    "complete": "#16a34a",
    "completed": "#16a34a",
    "at risk": "#d97706",
    "blocked": "#dc2626",
    "off track": "#dc2626",
}


def _status_color(status: str) -> str:
    return _STATUS_COLORS.get(status.strip().lower(), "#64748b")


def generate_html(report: MilestoneReport) -> str:
    metric_cards = "".join(
        f"""<article class="card">
          <span class="label">{escape(metric.name)}</span>
          <strong>{escape(metric.value)}</strong>
          <span>Target: {escape(metric.target or "—")}</span>
          <span class="status" style="--status:{_status_color(metric.status)}">{escape(metric.status)}</span>
        </article>"""
        for metric in report.metrics
    ) or '<p class="empty">No metrics were supplied.</p>'

    milestone_rows = "".join(
        f"""<tr>
          <td>{escape(item.name)}</td><td>{escape(item.owner)}</td>
          <td>{escape(item.due or "—")}</td>
          <td><span class="status" style="--status:{_status_color(item.status)}">{escape(item.status)}</span></td>
          <td><div class="progress" aria-label="{item.progress}% complete"><i style="width:{item.progress}%"></i></div>{item.progress}%</td>
        </tr>"""
        for item in report.milestones
    ) or '<tr><td colspan="5" class="empty">No milestones were supplied.</td></tr>'

    risks = "".join(f"<li>{escape(risk)}</li>" for risk in report.risks)
    risk_section = (
        f"<section><h2>Risks &amp; dependencies</h2><ul>{risks}</ul></section>"
        if risks
        else ""
    )

    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{escape(report.title)} dashboard</title>
  <style>
    :root {{ color-scheme: light; font-family: Inter, system-ui, sans-serif; color:#14233d; background:#f1f5f9 }}
    * {{ box-sizing:border-box }} body {{ margin:0 }} header {{ color:white; background:#14233d; padding:2.5rem max(5vw,1rem) }}
    header p {{ color:#cbd5e1; margin:.4rem 0 0 }} main {{ max-width:1200px; margin:auto; padding:2rem max(3vw,1rem) }}
    h1 {{ margin:0; font-size:clamp(2rem,4vw,3.25rem) }} h2 {{ margin:0 0 1rem }}
    .summary, section {{ background:white; border-radius:14px; padding:1.5rem; margin-bottom:1.5rem; box-shadow:0 2px 10px #0f172a0d }}
    .grid {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(190px,1fr)); gap:1rem; margin-bottom:1.5rem }}
    .card {{ background:white; border-radius:14px; padding:1.25rem; display:grid; gap:.45rem; box-shadow:0 2px 10px #0f172a0d }}
    .card strong {{ font-size:2rem }} .label {{ color:#475569; font-weight:650 }}
    .status {{ color:var(--status); font-weight:700 }} table {{ border-collapse:collapse; width:100% }}
    th,td {{ text-align:left; padding:.85rem .6rem; border-bottom:1px solid #e2e8f0 }} th {{ color:#475569 }}
    .progress {{ display:inline-block; width:90px; height:8px; margin-right:.5rem; background:#e2e8f0; border-radius:99px; overflow:hidden }}
    .progress i {{ display:block; height:100%; background:#2563eb }} .empty {{ color:#64748b }} ul {{ margin-bottom:0 }}
    @media(max-width:700px) {{ table {{ display:block; overflow-x:auto }} }}
    @media print {{ body {{ background:white }} .summary,section,.card {{ box-shadow:none; border:1px solid #e2e8f0 }} }}
  </style>
</head>
<body>
  <header><h1>{escape(report.title)}</h1><p>Reporting period: {escape(report.period)}</p></header>
  <main>
    <div class="summary"><h2>Executive summary</h2><p>{escape(report.summary or "No executive summary was supplied.")}</p></div>
    <h2>Key metrics</h2><div class="grid">{metric_cards}</div>
    <section><h2>Milestones</h2><table>
      <thead><tr><th>Milestone</th><th>Owner</th><th>Due</th><th>Status</th><th>Progress</th></tr></thead>
      <tbody>{milestone_rows}</tbody>
    </table></section>
    {risk_section}
  </main>
</body>
</html>
"""


def _add_title(slide, title: str, subtitle: str = "") -> None:
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(11.9), Inches(0.7))
    paragraph = title_box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True
    paragraph.font.color.rgb = _NAVY
    if subtitle:
        subtitle_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.15), Inches(11.9), Inches(0.4))
        subtitle_paragraph = subtitle_box.text_frame.paragraphs[0]
        subtitle_paragraph.text = subtitle
        subtitle_paragraph.font.size = Pt(12)
        subtitle_paragraph.font.color.rgb = RGBColor(100, 116, 139)


def generate_powerpoint(report: MilestoneReport, destination: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    cover = presentation.slides.add_slide(presentation.slide_layouts[6])
    cover.background.fill.solid()
    cover.background.fill.fore_color.rgb = _NAVY
    title_box = cover.shapes.add_textbox(Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.2))
    title = title_box.text_frame.paragraphs[0]
    title.text = report.title
    title.font.size = Pt(36)
    title.font.bold = True
    title.font.color.rgb = RGBColor(255, 255, 255)
    subtitle = cover.shapes.add_textbox(Inches(0.9), Inches(3.6), Inches(11), Inches(0.5))
    subtitle.text_frame.paragraphs[0].text = f"Milestone dashboard  •  {report.period}"
    subtitle.text_frame.paragraphs[0].font.color.rgb = RGBColor(203, 213, 225)

    overview = presentation.slides.add_slide(presentation.slide_layouts[6])
    _add_title(overview, "Executive overview", report.summary or "No executive summary was supplied.")
    metrics = report.metrics or ()
    for index, metric in enumerate(metrics[:8]):
        column, row = index % 4, index // 4
        left, top = 0.7 + column * 3.1, 1.9 + row * 2.3
        shape = overview.shapes.add_textbox(Inches(left), Inches(top), Inches(2.7), Inches(1.7))
        frame = shape.text_frame
        frame.text = metric.name
        frame.paragraphs[0].font.size = Pt(14)
        frame.paragraphs[0].font.bold = True
        value = frame.add_paragraph()
        value.text = metric.value
        value.font.size = Pt(25)
        value.font.bold = True
        value.font.color.rgb = _BLUE
        target = frame.add_paragraph()
        target.text = f"Target: {metric.target or '—'}  •  {metric.status}"
        target.font.size = Pt(11)
    if not metrics:
        empty = overview.shapes.add_textbox(Inches(0.7), Inches(2), Inches(11.9), Inches(0.5))
        empty.text_frame.paragraphs[0].text = "No metrics were supplied."

    milestones = presentation.slides.add_slide(presentation.slide_layouts[6])
    _add_title(milestones, "Milestone plan", f"{len(report.milestones)} milestones")
    rows = max(2, min(len(report.milestones), 10) + 1)
    table = milestones.shapes.add_table(rows, 5, Inches(0.7), Inches(1.65), Inches(11.9), Inches(4.9)).table
    widths = [3.4, 2.0, 1.8, 2.4, 2.3]
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for cell, heading in zip(table.rows[0].cells, ("Milestone", "Owner", "Due", "Status", "Progress")):
        cell.text = heading
        cell.fill.solid()
        cell.fill.fore_color.rgb = _NAVY
        cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        cell.text_frame.paragraphs[0].font.bold = True
    if report.milestones:
        for row_index, item in enumerate(report.milestones[:9], start=1):
            row = table.rows[row_index]
            for cell, value in zip(
                row.cells,
                (item.name, item.owner, item.due or "—", item.status, f"{item.progress}%"),
            ):
                cell.text = value
                cell.text_frame.paragraphs[0].font.size = Pt(11)
                cell.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    else:
        table.cell(1, 0).text = "No milestones were supplied."

    destination.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(destination)


def generate_dashboard(report: MilestoneReport, output_directory: Path | str) -> tuple[Path, Path]:
    output = Path(output_directory)
    output.mkdir(parents=True, exist_ok=True)
    html_path = output / "dashboard.html"
    powerpoint_path = output / "dashboard.pptx"
    html_path.write_text(generate_html(report), encoding="utf-8")
    generate_powerpoint(report, powerpoint_path)
    return html_path, powerpoint_path
