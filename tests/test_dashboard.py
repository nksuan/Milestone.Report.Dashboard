import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

from milestone_dashboard import generate_dashboard, parse_report
from milestone_dashboard.dashboard import generate_html


SAMPLE_PROMPT = """\
Title: Q3 Platform Migration
Period: July–September 2026
Summary: Migration is progressing to plan.
Metric: Services migrated | 18 | 24 | On track
Milestone: Complete pilot | Ada | 2026-08-15 | Completed | 100%
Milestone: Production cutover | Sam | 2026-09-20 | At risk | 65
Risk: Vendor access may delay the cutover.
"""


class ReportParserTests(unittest.TestCase):
    def test_parses_prompt_fields(self):
        report = parse_report(SAMPLE_PROMPT)

        self.assertEqual(report.title, "Q3 Platform Migration")
        self.assertEqual(report.metrics[0].value, "18")
        self.assertEqual(report.milestones[1].progress, 65)
        self.assertEqual(report.risks, ("Vendor access may delay the cutover.",))

    def test_rejects_empty_prompt(self):
        with self.assertRaisesRegex(ValueError, "cannot be empty"):
            parse_report(" \n ")

    def test_clamps_invalid_progress(self):
        report = parse_report("Milestone: Launch | Owner | Tomorrow | Blocked | 150")
        self.assertEqual(report.milestones[0].progress, 100)


class DashboardTests(unittest.TestCase):
    def test_html_is_self_contained_and_escapes_prompt_content(self):
        report = parse_report("Title: <Launch>\nSummary: Ready & waiting")
        html = generate_html(report)

        self.assertIn("&lt;Launch&gt;", html)
        self.assertIn("Ready &amp; waiting", html)
        self.assertNotIn("<Launch>", html)

    def test_generates_html_and_powerpoint(self):
        report = parse_report(SAMPLE_PROMPT)
        with tempfile.TemporaryDirectory() as directory:
            html_path, powerpoint_path = generate_dashboard(report, directory)

            self.assertTrue(html_path.is_file())
            self.assertTrue(powerpoint_path.is_file())
            presentation = Presentation(powerpoint_path)
            self.assertEqual(len(presentation.slides), 3)
            self.assertEqual(
                presentation.slides[0].shapes[0].text,
                "Q3 Platform Migration",
            )


if __name__ == "__main__":
    unittest.main()
